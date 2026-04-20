from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT = "/Users/kraigwastlund/Sources/UT/CS3005_Spring2026/Final Project/RobotWarz_Project_Lecture.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

COL_BG = RGBColor(245, 247, 250)
COL_PRIMARY = RGBColor(21, 76, 121)
COL_ACCENT = RGBColor(227, 114, 34)
COL_TEXT = RGBColor(28, 37, 46)
COL_MUTED = RGBColor(98, 110, 123)
COL_WHITE = RGBColor(255, 255, 255)
COL_GREEN = RGBColor(36, 123, 83)
COL_RED = RGBColor(170, 34, 34)


def set_bg(slide, color=COL_BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(9.2), Inches(1.0))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = COL_PRIMARY

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.72), Inches(1.2), Inches(11.8), Inches(0.7))
        stf = sub.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        srun = sp.add_run()
        srun.text = subtitle
        srun.font.size = Pt(18)
        srun.font.color.rgb = COL_MUTED


def add_footer(slide, text="CS 3005 - RobotWarz Final Project"):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.15), Inches(13.333), Inches(0.35))
    line.fill.solid()
    line.fill.fore_color.rgb = COL_PRIMARY
    line.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.17), Inches(12.2), Inches(0.25))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(11)
    run.font.color.rgb = COL_WHITE


def add_bullets(slide, x, y, w, h, bullets, level0_size=24, leveln_size=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(bullets):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(level0_size if level == 0 else leveln_size)
        p.font.color.rgb = COL_TEXT


def two_col_boxes(slide, left_title, left_items, right_title, right_items):
    lx, ly, lw, lh = Inches(0.8), Inches(1.7), Inches(5.8), Inches(4.9)
    rx, ry, rw, rh = Inches(6.75), Inches(1.7), Inches(5.8), Inches(4.9)

    for x, y, w, h, title, items in [
        (lx, ly, lw, lh, left_title, left_items),
        (rx, ry, rw, rh, right_title, right_items),
    ]:
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        rect.fill.solid()
        rect.fill.fore_color.rgb = COL_WHITE
        rect.line.color.rgb = RGBColor(220, 226, 232)
        tbox = slide.shapes.add_textbox(x + Inches(0.25), y + Inches(0.18), w - Inches(0.5), Inches(0.5))
        ttf = tbox.text_frame
        p = ttf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = COL_PRIMARY
        add_bullets(slide, (x.inches + 0.25), (y.inches + 0.8), (w.inches - 0.5), (h.inches - 1.0), items, 18, 16)


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "RobotWarz Final Project", "CS 3005 | Build an arena. Design a robot. Make them battle.")
hero = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(2.0), Inches(11.8), Inches(3.9))
hero.fill.solid()
hero.fill.fore_color.rgb = COL_WHITE
hero.line.color.rgb = RGBColor(218, 224, 230)
add_bullets(slide, 1.2, 2.35, 11.0, 3.2, [
    "In-class final: Monday, May 4, 11:00-12:50, Smith 117",
    "Take-home goal: functional Robot Battle simulation",
    "You build the Arena and at least one custom robot",
    "Starter code gives RobotBase, RadarObj, samples, and test harness",
], 24, 20)
add_footer(slide)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Learning Goals")
two_col_boxes(
    slide,
    "Technical Goals",
    [
        "Practice OOP with polymorphism and shared interfaces",
        "Build a turn-based simulation engine",
        "Use dynamic loading (.so) at runtime",
        "Model game rules with clean, testable logic",
    ],
    "Engineering Goals",
    [
        "Work modularly: loading, movement, shooting, UI",
        "Iterate from simple to robust",
        "Follow strict interface contracts",
        "Write portable robots that work in other arenas",
    ],
)
add_footer(slide)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "High-Level Architecture")
arch = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.8), Inches(11.6), Inches(4.8))
arch.fill.solid(); arch.fill.fore_color.rgb = COL_WHITE; arch.line.color.rgb = RGBColor(220,226,232)

boxes = [
    (1.3, 2.5, 2.7, 1.2, "Robot_<Name>.cpp", COL_ACCENT),
    (4.4, 2.5, 2.7, 1.2, "libRobot.so", COL_PRIMARY),
    (7.5, 2.5, 3.0, 1.2, "Arena Simulation", COL_GREEN),
]
for x, y, w, h, label, color in boxes:
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = color; b.line.fill.background()
    tf = b.text_frame; tf.clear(); p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = COL_WHITE

ar1 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.0), Inches(2.88), Inches(0.35), Inches(0.45))
ar1.fill.solid(); ar1.fill.fore_color.rgb = COL_MUTED; ar1.line.fill.background()
ar2 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(7.1), Inches(2.88), Inches(0.35), Inches(0.45))
ar2.fill.solid(); ar2.fill.fore_color.rgb = COL_MUTED; ar2.line.fill.background()

add_bullets(slide, 1.3, 4.2, 10.9, 1.9, [
    "Arena compiles each Robot_<Name>.cpp into a shared library (.so)",
    "Arena loads create_robot() and robot_summary() at runtime",
    "Arena drives turns through RobotBase virtual functions",
], 18, 16)
add_footer(slide)

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Robot Contract (Non-Negotiable)")
panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.6), Inches(11.6), Inches(5.0))
panel.fill.solid(); panel.fill.fore_color.rgb = COL_WHITE; panel.line.color.rgb = RGBColor(220,226,232)

code = slide.shapes.add_textbox(Inches(1.3), Inches(2.1), Inches(10.9), Inches(1.4))
ctf = code.text_frame
ctf.word_wrap = True
p = ctf.paragraphs[0]
r = p.add_run(); r.text = 'extern "C" RobotBase* create_robot();\nextern "C" const char* robot_summary();'
r.font.name = "Courier New"; r.font.size = Pt(24); r.font.color.rgb = COL_PRIMARY; r.font.bold = True

add_bullets(slide, 1.3, 3.7, 10.9, 2.5, [
    "robot_summary() is required for grading compatibility",
    "Summary must be 1 to 50 characters",
    "Robots must derive from RobotBase and implement required methods",
    "Do NOT modify RobotBase.h, RobotBase.cpp, or RadarObj.h",
], 20, 18)
add_footer(slide)

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Turn Flow")
two_col_boxes(
    slide,
    "Per Robot Turn",
    [
        "1. Arena asks for radar direction",
        "2. Arena scans and sends vector<RadarObj>",
        "3. Robot chooses one action",
        "4. Arena applies legal move or shot",
        "5. Arena updates health, armor, death state",
    ],
    "Action Rule",
    [
        "Exactly one action per turn:",
        ("Shoot OR Move OR Do Nothing", 1),
        "Never both shoot and move in one turn",
        "Dead robots do not take turns",
    ],
)
add_footer(slide)

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Weapons and Obstacles")
two_col_boxes(
    slide,
    "Weapons",
    [
        "Flamethrower: 3x4 area",
        "Railgun: line across arena, pierces",
        "Hammer: one adjacent cell",
        "Grenade: target cell, 3x3 splash, 10 shots max",
        "Any hit reduces armor by 1",
    ],
    "Obstacles",
    [
        "F: flamethrower obstacle (damages robot)",
        "P: pit (move speed becomes 0)",
        "M: mound (blocks movement)",
        "Mounds do not block shots",
    ],
)
add_footer(slide)

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Suggested Build Order")
add_bullets(slide, 1.0, 1.9, 11.4, 4.8, [
    "Milestone 1: Board + robot placement + board printing",
    "Milestone 2: Dynamic robot compile/load (.so)",
    "Milestone 3: Radar plumbing with RadarObj",
    "Milestone 4: Movement legality (bounds + obstacles)",
    "Milestone 5: Shooting and damage resolution",
    "Milestone 6: Round loop, winner detection, clean output",
    "Milestone 7: polish UI, edge cases, and stress tests",
], 24, 20)
add_footer(slide)

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Testing Workflow")
two_col_boxes(
    slide,
    "Robot-Level Testing",
    [
        "make",
        "./test_robot Robot_MyBot.cpp",
        "Verify required exports and summary length",
        "Validate action behavior under radar input",
    ],
    "Arena-Level Testing",
    [
        "Run full simulation with multiple robots",
        "Test collisions, pits, mounds, and obstacle damage",
        "Confirm no segfaults/endless loops",
        "Ensure one winner is always produced",
    ],
)
add_footer(slide)

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "What I Will Grade")
box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(12.0), Inches(4.9))
box.fill.solid(); box.fill.fore_color.rgb = COL_WHITE; box.line.color.rgb = RGBColor(220,226,232)

add_bullets(slide, 1.15, 2.1, 11.3, 3.9, [
    "0 if project does not build/run from a fresh clone",
    "10 - basic structure and spec adherence",
    "10 - working simulation and robot interaction",
    "10 - legal movement and boundary handling",
    "20 - shooting and damage logic",
    "10 - clear output/UI",
    "10 - game resolution (last robot alive)",
    "Deductions for crashes, invalid behaviors, endless loops",
], 22, 18)
add_footer(slide)

# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Common Failure Modes to Avoid")
warning = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(0.9), Inches(2.0), Inches(1.1), Inches(1.0))
warning.fill.solid(); warning.fill.fore_color.rgb = COL_RED; warning.line.fill.background()
add_bullets(slide, 2.3, 1.9, 10.2, 4.9, [
    "Changing RobotBase or RadarObj files",
    "Robots that rely on arena internals",
    "Allowing move + shoot in same turn",
    "Ignoring weapon range/shape rules",
    "Illegal movement through mounds or out-of-bounds",
    "No clear game-over condition",
    "Not testing from clean build",
], 24, 20)
add_footer(slide)

# Slide 11
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_title(slide, "Student Deliverables Checklist")
check = [
    "RobotWarz executable builds and runs with make",
    "Arena implementation handles full game loop",
    "At least one custom Robot_<Name>.cpp",
    "Robot exports create_robot() and robot_summary()",
    "robot_summary length is 1-50 chars",
    "No changes to RobotBase/RadarObj starter contracts",
    "Readable board output + turn narration",
]
for i, txt in enumerate(check):
    y = 1.9 + i * 0.65
    sq = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(y), Inches(0.35), Inches(0.35))
    sq.fill.solid(); sq.fill.fore_color.rgb = COL_GREEN; sq.line.fill.background()
    t = slide.shapes.add_textbox(Inches(1.45), Inches(y - 0.03), Inches(10.8), Inches(0.45))
    p = t.text_frame.paragraphs[0]
    r = p.add_run(); r.text = txt; r.font.size = Pt(22); r.font.color.rgb = COL_TEXT
add_footer(slide)

# Slide 12
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, RGBColor(21, 76, 121))
add_title(slide, "You Can Absolutely Build This", "Start simple. Iterate. Test often. Ask questions early.")
for shp in slide.shapes:
    if shp.has_text_frame:
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                if r.font.color is not None:
                    r.font.color.rgb = COL_WHITE

mid = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.4), Inches(2.4))
p = mid.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Project success formula:\nBuild core loop -> add rules -> polish UI -> harden edge cases"
r.font.size = Pt(32)
r.font.bold = True
r.font.color.rgb = COL_WHITE

add_footer(slide, "CS 3005 - RobotWarz")

prs.save(OUT)
print(OUT)
